from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from collector.adapters.excel_laliga import _SQUADS_HEADERS, _TRANSFERS_HEADERS

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "laliga-fantasy-oficial.xlsx"

NAVY = "2F4A89"
GOLD = "F4A32E"
INK = "1A1A1A"
MUTED = "5C5C5C"
LINE = "D0D5DD"
EXAMPLE = "FFF4D6"
HEADER = "2F4A89"
WHITE = "FFFFFF"

SQUAD_EXAMPLE = [
    [1, "GW1", "finished", 64, False, 0, "EXAMPLE — delete this row", "FWD", "Real Madrid", "starter", True, False, 12, 7.4, 15.0, "", ""],
    [1, "GW1", "finished", 64, False, 0, "EXAMPLE — bench placeholder", "MID", "Barcelona", "bench", False, True, 2, 6.2, 8.0, "", ""],
]

TRANSFER_EXAMPLE = [
    [2, "GW2", "EXAMPLE — player in", "Real Madrid", "FWD", 15.0, "", "", "EXAMPLE — player out", "Barcelona", "MID", 8.0, "", "", 0],
]


def build(path: Path = OUTPUT) -> Path:
    book = Workbook()
    _readme(book.active)
    _competition(book.create_sheet("Competition"))
    squads = book.create_sheet("Squads")
    transfers = book.create_sheet("Transfers")
    _clubs(book.create_sheet("Clubs"))
    _squads(squads)
    _transfers(transfers)
    book.save(path)
    return path


def _readme(sheet: Worksheet) -> None:
    sheet.title = "Read me"
    sheet.sheet_properties.tabColor = GOLD
    sheet["A1"] = "Official LaLiga Fantasy — weekly workbook"
    sheet["A1"].font = Font(name="Calibri", size=18, bold=True, color=NAVY)
    lines = [
        "",
        "The official app has no website JSON export. Copy one week at a time from the Android app into this file.",
        "Keep a working copy (for example in Documents). Leave this template in collector/templates/ untouched.",
        "",
        "1. Competition — your team name, manager, and season (2026/27).",
        "2. Squads — one row per player, every gameweek. Repeat teamPoints / tripleCaptain / transferPenalty on each row of that week.",
        "3. Transfers — optional. One row per in/out pair. Skip a week if you made no transfers.",
        "4. Clubs — SofaScore club IDs for portraits later. Copy sofaScoreClubId onto Squads/Transfers if you want crests now.",
        "",
        "Rules",
        "• Do not rename sheets or header cells.",
        "• Delete the yellow EXAMPLE rows before import.",
        "• teamPoints is the app’s week total (already includes captain / chips). Do not double it.",
        "• Captain: TRUE on one starter. Triple captain: TRUE on every row of that week if you played the chip.",
        "• Positions: GK, DEF, MID, FWD. Roles: starter or bench.",
        "• sofaScorePlayerId / sofaScoreClubId are optional. Blank is fine; we can match names later.",
        "",
        "Import (from the collector folder, backend running):",
        "  python -m collector import-excel templates\\laliga-fantasy-oficial.xlsx --dry-run",
        "  python -m collector import-excel path\\to\\your-copy.xlsx",
        "  python -m collector import-excel path\\to\\your-copy.xlsx --gameweek 3",
        "",
        "This is a separate competition from SofaScore LaLiga in the dashboard.",
    ]
    for index, line in enumerate(lines, start=2):
        sheet[f"A{index}"] = line
        sheet[f"A{index}"].font = Font(name="Calibri", size=12, color=INK, bold=line.startswith("Rules"))
        sheet[f"A{index}"].alignment = Alignment(wrap_text=True)
    sheet.column_dimensions["A"].width = 118
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"


def _competition(sheet: Worksheet) -> None:
    sheet.sheet_properties.tabColor = NAVY
    sheet["A1"] = "Field"
    sheet["B1"] = "Value"
    _header_row(sheet, 1, 2)
    fields = [
        ("teamName", ""),
        ("managerName", "Locksat"),
        ("season", "2026/27"),
        ("competitionName", "LaLiga Fantasy"),
    ]
    notes = {
        "teamName": "Required. The name shown in the official app.",
        "managerName": "Optional.",
        "season": "Use 2026/27 unless you start a later season.",
        "competitionName": "Shown in the dashboard dropdown.",
    }
    for index, (field, value) in enumerate(fields, start=2):
        sheet[f"A{index}"] = field
        sheet[f"B{index}"] = value
        sheet[f"C{index}"] = notes[field]
        sheet[f"A{index}"].font = Font(bold=True, color=NAVY)
        sheet[f"C{index}"].font = Font(italic=True, color=MUTED, size=11)
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 36
    sheet.column_dimensions["C"].width = 56
    sheet.freeze_panes = "A2"


def _squads(sheet: Worksheet) -> None:
    sheet.sheet_properties.tabColor = GOLD
    sheet.append(list(_SQUADS_HEADERS))
    _header_row(sheet, 1, len(_SQUADS_HEADERS))
    for row in SQUAD_EXAMPLE:
        sheet.append(row)
    for index in range(3, 5):
        _fill_row(sheet, index, len(_SQUADS_HEADERS), EXAMPLE)
    for _ in range(20):
        sheet.append([None] * len(_SQUADS_HEADERS))
    widths = [12, 14, 12, 13, 15, 16, 28, 12, 22, 12, 12, 13, 10, 10, 10, 20, 18]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(_SQUADS_HEADERS))}25"
    _list(sheet, "C2:C200", "finished,live,upcoming")
    _list(sheet, "H2:H200", "GK,DEF,MID,FWD")
    _list(sheet, "J2:J200", "starter,bench")
    _list(sheet, "E2:E200", "TRUE,FALSE")
    _list(sheet, "K2:K200", "TRUE,FALSE")
    _list(sheet, "L2:L200", "TRUE,FALSE")


def _transfers(sheet: Worksheet) -> None:
    sheet.sheet_properties.tabColor = "8A8A8A"
    sheet.append(list(_TRANSFERS_HEADERS))
    _header_row(sheet, 1, len(_TRANSFERS_HEADERS))
    for row in TRANSFER_EXAMPLE:
        sheet.append(row)
    _fill_row(sheet, 2, len(_TRANSFERS_HEADERS), EXAMPLE)
    for _ in range(15):
        sheet.append([None] * len(_TRANSFERS_HEADERS))
    for index in range(1, len(_TRANSFERS_HEADERS) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 18
    sheet.column_dimensions["C"].width = 28
    sheet.column_dimensions["I"].width = 28
    sheet.freeze_panes = "A2"
    _list(sheet, "E2:E200", "GK,DEF,MID,FWD")
    _list(sheet, "K2:K200", "GK,DEF,MID,FWD")


def _clubs(sheet: Worksheet) -> None:
    sheet.sheet_properties.tabColor = "1D9E75"
    sheet["A1"] = "club"
    sheet["B1"] = "sofaScoreClubId"
    sheet["C1"] = "notes"
    _header_row(sheet, 1, 3)
    clubs = [
        ("Athletic Club", "2825", ""),
        ("Atlético Madrid", "2836", ""),
        ("Barcelona", "2817", "FC Barcelona"),
        ("Celta Vigo", "2821", ""),
        ("Deportivo Alavés", "2885", "Alavés"),
        ("Deportivo de La Coruña", "2832", "Deportivo"),
        ("Elche", "2846", ""),
        ("Espanyol", "2814", ""),
        ("Getafe", "2859", ""),
        ("Levante", "2849", "Levante UD"),
        ("Málaga", "2830", "Málaga CF"),
        ("Osasuna", "2820", ""),
        ("Rayo Vallecano", "2818", ""),
        ("Real Betis", "2816", ""),
        ("Real Madrid", "2829", ""),
        ("Real Racing Club", "2835", "Racing Santander"),
        ("Real Sociedad", "2824", ""),
        ("Sevilla", "2833", ""),
        ("Valencia", "2828", ""),
        ("Villarreal", "2819", ""),
    ]
    for row in clubs:
        sheet.append(list(row))
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 28
    sheet.freeze_panes = "A2"
    note = sheet["A23"]
    note.value = "IDs are SofaScore LaLiga club IDs (unique tournament 8). Confirm in the app vs dashboard if a newly promoted club looks wrong. Player portraits need sofaScorePlayerId on Squads."
    note.font = Font(italic=True, color=MUTED)
    sheet.merge_cells("A23:C23")


def _header_row(sheet: Worksheet, row: int, columns: int) -> None:
    fill = PatternFill("solid", fgColor=HEADER)
    font = Font(bold=True, color=WHITE)
    thin = Border(bottom=Side(style="thin", color=LINE))
    for index in range(1, columns + 1):
        cell = sheet.cell(row, index)
        cell.fill = fill
        cell.font = font
        cell.border = thin
        cell.alignment = Alignment(horizontal="center")


def _fill_row(sheet: Worksheet, row: int, columns: int, color: str) -> None:
    fill = PatternFill("solid", fgColor=color)
    for index in range(1, columns + 1):
        sheet.cell(row, index).fill = fill


def _list(sheet: Worksheet, ref: str, formula: str) -> None:
    validation = DataValidation(type="list", formula1=f'"{formula}"', allow_blank=True)
    validation.error = "Pick a value from the list"
    validation.errorTitle = "Invalid value"
    sheet.add_data_validation(validation)
    validation.add(ref)


if __name__ == "__main__":
    target = build()
    print(target)
