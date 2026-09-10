from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from collector.models import (
    PickPayload,
    PlayerPayload,
    Snapshot,
    TransferPair,
    TransferPlayer,
    TransferRound,
    TransfersBatch,
)

SOURCE = "laliga-fantasy"
EXTERNAL_ID = "oficial"
DEFAULT_NAME = "LaLiga Fantasy"
DEFAULT_SLUG = "laliga-fantasy-oficial"
DEFAULT_SEASON = "2026/27"

_SQUADS_HEADERS = (
    "gameweek",
    "gameweekName",
    "status",
    "teamPoints",
    "playerName",
    "position",
    "club",
    "role",
    "points",
    "rating",
    "price",
    "sofaScorePlayerId",
    "sofaScoreClubId",
)

_TRANSFERS_HEADERS = (
    "gameweek",
    "gameweekName",
    "action",
    "playerName",
    "club",
    "position",
    "price",
    "counterpart",
    "sofaScorePlayerId",
    "sofaScoreClubId",
)


def load_workbook_payloads(
    path: str | Path,
    *,
    gameweek: int | None = None,
) -> tuple[list[Snapshot], TransfersBatch | None]:
    """Read the official LaLiga Fantasy template into ingest snapshots."""
    workbook = load_workbook(Path(path), data_only=True, read_only=True)
    try:
        competition, team = _competition_and_team(workbook)
        snapshots = _snapshots_from_squads(workbook, competition, team, gameweek)
        batch = _transfers_from_sheet(workbook, competition, team, gameweek)
    finally:
        workbook.close()
    if not snapshots:
        raise ValueError(
            "No squad rows to import. Add players on the Squads sheet and delete the EXAMPLE rows."
        )
    return snapshots, batch


def _competition_and_team(workbook: Any) -> tuple[dict[str, Any], dict[str, Any]]:
    sheet = workbook["Competition"]
    values = _key_values(sheet)
    team_name = (values.get("teamName") or "").strip() or (values.get("managerName") or "").strip()
    if not team_name:
        raise ValueError(
            "Competition.teamName is required, or set managerName if the official app has no team name."
        )
    season = values.get("season") or DEFAULT_SEASON
    name = values.get("competitionName") or DEFAULT_NAME
    return (
        {
            "source": SOURCE,
            "externalId": EXTERNAL_ID,
            "name": name,
            "season": season,
            "slug": DEFAULT_SLUG,
        },
        {
            "name": team_name,
            "managerName": values.get("managerName") or None,
        },
    )


def _snapshots_from_squads(
    workbook: Any,
    competition: dict[str, Any],
    team: dict[str, Any],
    gameweek: int | None,
) -> list[Snapshot]:
    rows = _table_rows(workbook["Squads"], _SQUADS_HEADERS)
    grouped: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        name = str(row.get("playerName") or "").strip()
        if not name or name.upper().startswith("EXAMPLE"):
            continue
        number = _int(row.get("gameweek"), "Squads.gameweek")
        if gameweek is not None and number != gameweek:
            continue
        grouped[number].append(row)
    snapshots: list[Snapshot] = []
    for number, items in sorted(grouped.items()):
        first = items[0]
        picks = [_pick_from_row(item) for item in items]
        snapshots.append(
            Snapshot(
                competition=competition,
                gameweek={
                    "number": number,
                    "name": str(first.get("gameweekName") or f"GW{number}"),
                    "status": str(first.get("status") or "finished").strip().lower(),
                },
                team=team,
                picks=picks,
                teamPoints=_optional_float(first.get("teamPoints")),
                tripleCaptain=False,
                transferPenalty=None,
            )
        )
    return snapshots


def _transfers_from_sheet(
    workbook: Any,
    competition: dict[str, Any],
    team: dict[str, Any],
    gameweek: int | None,
) -> TransfersBatch | None:
    if "Transfers" not in workbook.sheetnames:
        return None
    rows = _table_rows(workbook["Transfers"], _TRANSFERS_HEADERS)
    grouped: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        name = str(row.get("playerName") or "").strip()
        if not name or name.upper().startswith("EXAMPLE"):
            continue
        number = _int(row.get("gameweek"), "Transfers.gameweek")
        if gameweek is not None and number != gameweek:
            continue
        grouped[number].append(row)
    if not grouped:
        return None
    rounds: list[TransferRound] = []
    for number, items in sorted(grouped.items()):
        first = items[0]
        pairs = [_deal_from_row(item) for item in items]
        rounds.append(
            TransferRound(
                number=number,
                name=str(first.get("gameweekName") or f"GW{number}"),
                transferPenalty=0,
                transfers=pairs,
            )
        )
    return TransfersBatch(competition=competition, team=team, rounds=rounds)


def _pick_from_row(row: dict[str, Any]) -> PickPayload:
    name = str(row.get("playerName") or "").strip()
    player_id = _id_or_slug(row.get("sofaScorePlayerId"), name, row.get("club"))
    club_id = _digits(row.get("sofaScoreClubId"))
    return PickPayload(
        player=PlayerPayload(
            externalId=player_id,
            name=name,
            position=_position(row.get("position")),
            club=_optional_text(row.get("club")),
            clubExternalId=club_id,
        ),
        role=_role(row.get("role")),
        points=_optional_float(row.get("points")) or 0,
        captain=False,
        viceCaptain=False,
        rating=_optional_float(row.get("rating")),
        price=_optional_float(row.get("price")),
    )


def _deal_from_row(row: dict[str, Any]) -> TransferPair:
    player = _transfer_player(
        row.get("playerName"),
        row.get("club"),
        row.get("position"),
        row.get("price"),
        row.get("sofaScorePlayerId"),
        row.get("sofaScoreClubId"),
    )
    counterpart = _optional_text(row.get("counterpart")) or "Market"
    action = _action(row.get("action"))
    if action == "sold":
        return TransferPair(playerOut=player, counterpart=counterpart)
    return TransferPair(playerIn=player, counterpart=counterpart)


def _transfer_player(
    name: Any,
    club: Any,
    position: Any,
    price: Any,
    player_id: Any,
    club_id: Any,
) -> TransferPlayer:
    label = str(name or "").strip()
    return TransferPlayer(
        externalId=_id_or_slug(player_id, label, club),
        name=label,
        position=_position(position),
        club=_optional_text(club),
        clubExternalId=_digits(club_id),
        price=_optional_float(price),
    )


def _key_values(sheet: Any) -> dict[str, str]:
    values: dict[str, str] = {}
    for row in sheet.iter_rows(min_row=1, max_col=2, values_only=True):
        key = str(row[0] or "").strip()
        if not key or key.startswith("#") or key in {"Field", "field"}:
            continue
        value = row[1]
        if value is None or str(value).strip() == "":
            continue
        values[key] = str(value).strip()
    return values


def _table_rows(sheet: Any, headers: tuple[str, ...]) -> list[dict[str, Any]]:
    rows = list(sheet.iter_rows(min_row=1, values_only=True))
    if not rows:
        return []
    found = [str(cell or "").strip() for cell in rows[0]]
    missing = [name for name in headers if name not in found]
    if missing:
        raise ValueError(f"{sheet.title} is missing columns: {', '.join(missing)}")
    index = {name: found.index(name) for name in headers}
    payload: list[dict[str, Any]] = []
    for raw in rows[1:]:
        item = {name: raw[index[name]] if index[name] < len(raw) else None for name in headers}
        payload.append(item)
    return payload


def _id_or_slug(explicit: Any, name: str, club: Any) -> str:
    digits = _digits(explicit)
    if digits:
        return digits
    club_part = _slug(str(club or ""))
    name_part = _slug(name)
    if not name_part:
        raise ValueError("Player name is required")
    return f"{name_part}-{club_part}" if club_part else name_part


def _slug(value: str) -> str:
    text = value.strip().lower()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-")


def _digits(value: Any) -> str | None:
    if value is None or str(value).strip() == "":
        return None
    text = str(value).strip()
    if text.endswith(".0") and text.replace(".", "", 1).isdigit():
        text = text[:-2]
    if text.isdigit():
        return text
    return None


def _int(value: Any, field: str) -> int:
    if value is None or str(value).strip() == "":
        raise ValueError(f"{field} is required")
    return int(float(value))


def _optional_float(value: Any) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    return float(value)


def _optional_text(value: Any) -> str | None:
    text = str(value or "").strip()
    return text or None


def _position(value: Any) -> str | None:
    text = str(value or "").strip().upper()
    if not text:
        return None
    aliases = {"G": "GK", "GOALKEEPER": "GK", "D": "DEF", "DEFENDER": "DEF", "M": "MID", "MIDFIELDER": "MID", "F": "FWD", "FORWARD": "FWD"}
    return aliases.get(text, text)


def _action(value: Any) -> str:
    text = str(value or "bought").strip().lower()
    if text in {"sold", "sell", "out", "sale"}:
        return "sold"
    if text in {"bought", "buy", "in", "purchase", "purchased"}:
        return "bought"
    raise ValueError(f"Transfers.action must be Bought or Sold, not {value!r}")


def _role(value: Any) -> str:
    text = str(value or "starter").strip().lower()
    if text in {"squad", "bench", "sub", "substitute", "reserve"}:
        return "squad"
    return "starter"
