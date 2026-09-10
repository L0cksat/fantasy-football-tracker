from pathlib import Path

from openpyxl import Workbook

from collector.adapters.excel_laliga import (
    _SQUADS_HEADERS,
    _TRANSFERS_HEADERS,
    load_workbook_payloads,
)


def _competition(book: Workbook, team_name: str = "Los Oficiales") -> None:
    sheet = book.create_sheet("Competition")
    sheet["A2"] = "teamName"
    sheet["B2"] = team_name
    sheet["A3"] = "managerName"
    sheet["B3"] = "Locksat"
    sheet["A4"] = "season"
    sheet["B4"] = "2026/27"


def test_example_rows_are_skipped(tmp_path: Path):
    path = tmp_path / "examples.xlsx"
    book = Workbook()
    book.active.title = "Read me"
    _competition(book)
    squads = book.create_sheet("Squads")
    squads.append(list(_SQUADS_HEADERS))
    squads.append(
        [1, "GW1", "finished", 10, "EXAMPLE — delete this row", "FWD", "Sevilla", "starter", 4, None, None, "", ""]
    )
    book.save(path)
    try:
        load_workbook_payloads(path)
        raise AssertionError("example-only workbook should not import")
    except ValueError as exc:
        assert "No squad rows" in str(exc)


def test_import_excel_maps_squad_and_transfers(tmp_path: Path):
    path = tmp_path / "week.xlsx"
    book = Workbook()
    book.active.title = "Read me"
    _competition(book)
    squads = book.create_sheet("Squads")
    squads.append(list(_SQUADS_HEADERS))
    squads.append(
        [1, "GW1", "finished", 71, "EXAMPLE — skip", "FWD", "Real Madrid", "starter", 9, 7, 12, "", ""]
    )
    squads.append(
        [1, "GW1", "finished", 71, "Kylian Mbappé", "FWD", "Real Madrid", "starter", 14, 8.1, 20, 823944, 2829]
    )
    squads.append(
        [1, "GW1", "finished", 71, "Pedri", "MID", "Barcelona", "squad", "", "", 12, "", ""]
    )
    transfers = book.create_sheet("Transfers")
    transfers.append(list(_TRANSFERS_HEADERS))
    transfers.append(
        [2, "GW2", "Bought", "Lamine Yamal", "Barcelona", "FWD", 18, "Market", 1399376, 2817]
    )
    transfers.append(
        [2, "GW2", "Sold", "Joselu", "Real Madrid", "FWD", 9, "Otro Manager FC", "", 2829]
    )
    book.save(path)

    snapshots, batch = load_workbook_payloads(path)
    assert len(snapshots) == 1
    snapshot = snapshots[0]
    assert snapshot.competition["source"] == "laliga-fantasy"
    assert snapshot.competition["slug"] == "laliga-fantasy-oficial"
    assert snapshot.team["name"] == "Los Oficiales"
    assert snapshot.gameweek["number"] == 1
    assert snapshot.teamPoints == 71
    assert snapshot.tripleCaptain is False
    assert [pick.player.name for pick in snapshot.picks] == ["Kylian Mbappé", "Pedri"]
    assert snapshot.picks[0].captain is False
    assert snapshot.picks[0].role == "starter"
    assert snapshot.picks[0].player.externalId == "823944"
    assert snapshot.picks[0].player.clubExternalId == "2829"
    assert snapshot.picks[1].role == "squad"
    assert snapshot.picks[1].player.externalId == "pedri-barcelona"
    assert batch is not None
    assert batch.rounds[0].number == 2
    bought, sold = batch.rounds[0].transfers
    assert bought.playerIn is not None
    assert bought.playerOut is None
    assert bought.playerIn.name == "Lamine Yamal"
    assert bought.playerIn.price == 18
    assert bought.counterpart == "Market"
    assert sold.playerOut is not None
    assert sold.playerIn is None
    assert sold.playerOut.name == "Joselu"
    assert sold.playerOut.price == 9
    assert sold.counterpart == "Otro Manager FC"


def test_import_excel_can_filter_gameweek(tmp_path: Path):
    path = tmp_path / "weeks.xlsx"
    book = Workbook()
    book.active.title = "Read me"
    _competition(book)
    squads = book.create_sheet("Squads")
    squads.append(list(_SQUADS_HEADERS))
    squads.append([1, "GW1", "finished", 10, "A", "FWD", "Sevilla", "starter", 4, None, None, "", ""])
    squads.append([2, "GW2", "finished", 20, "B", "MID", "Sevilla", "starter", 6, None, None, "", ""])
    book.save(path)

    snapshots, batch = load_workbook_payloads(path, gameweek=2)
    assert [item.gameweek["number"] for item in snapshots] == [2]
    assert snapshots[0].teamPoints == 20
    assert batch is None


def test_blank_team_name_falls_back_to_manager(tmp_path: Path):
    path = tmp_path / "manager-only.xlsx"
    book = Workbook()
    book.active.title = "Read me"
    _competition(book, team_name="")
    squads = book.create_sheet("Squads")
    squads.append(list(_SQUADS_HEADERS))
    squads.append([1, "GW1", "finished", 10, "A", "FWD", "Sevilla", "starter", 4, None, None, "", ""])
    book.save(path)

    snapshots, _batch = load_workbook_payloads(path)
    assert snapshots[0].team["name"] == "Locksat"
    assert snapshots[0].team["managerName"] == "Locksat"
